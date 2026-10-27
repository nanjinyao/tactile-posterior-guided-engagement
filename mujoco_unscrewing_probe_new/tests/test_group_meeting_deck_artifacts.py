from __future__ import annotations

import hashlib
import json
from pathlib import Path
import re
import subprocess
import unittest

from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
TRACE = ROOT.parent / "group_meeting_trace"
STEM = "unscrewing_contact_state_pipeline_slides_0709"
MARKDOWN = TRACE / f"{STEM}.md"
PPTX = TRACE / f"{STEM}.pptx"
PDF = TRACE / f"{STEM}.pdf"
MANIFEST = TRACE / f"{STEM}_build_manifest.json"


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _markdown_slide_count(text: str) -> int:
    return len(text.split("\n---\n")) - 1


@unittest.skipUnless(MANIFEST.is_file(), "group-meeting deck not built")
class TestGroupMeetingDeckArtifacts(unittest.TestCase):
    def setUp(self) -> None:
        self.manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
        markdown_record = self.manifest.get("artifacts", {}).get("markdown", {})
        current_sha = _sha256(MARKDOWN)
        if markdown_record.get("sha256") != current_sha:
            self.skipTest(
                "PPT/PDF manifest is a historical snapshot of an older Markdown SHA"
            )

    def test_markdown_pptx_pdf_are_synchronized_and_self_contained(self) -> None:
        text = MARKDOWN.read_text(encoding="utf-8")
        slide_count = _markdown_slide_count(text)
        self.assertEqual(len(Presentation(str(PPTX)).slides), slide_count)
        pages = subprocess.check_output(["pdfinfo", str(PDF)], text=True)
        self.assertRegex(pages, rf"(?m)^Pages:\s+{slide_count}$")

        references = re.findall(r'(?:src|href)="([^"]+)"', text)
        references.extend(re.findall(r"\]\(([^)]+)\)", text))
        local = {
            (MARKDOWN.parent / value.split("#", 1)[0]).resolve()
            for value in references
            if value and not re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*:", value)
        }
        self.assertTrue(local)
        self.assertEqual([path for path in local if not path.is_file()], [])

        self.assertIn("主线 30 页", text)
        self.assertIn("以下 4 页仅作问答 backup", text)
        self.assertIn("comparison_fuzzy_retry3_2x2_poster.png", text)
        self.assertIn("comparison_clear_retry3_2x2_poster.png", text)
        self.assertIn("comparison_improved4_vs_jin3_retry3.mp4", text)
        self.assertIn("m8_strict_entropy_grid_topview_sobol_003.mp4", text)
        self.assertIn("same_start_shortcut_residual_sobol_008.png", text)
        self.assertIn(
            "cross_size_s3_full_loop_pilot_v3/cross_size_s3_full_loop_pilot.png",
            text,
        )
        self.assertIn("答案是 mixed", text)
        self.assertIn("扶正不是 terminal", text)
        self.assertIn("R3 提高可恢复扰动完成率", text)
        self.assertIn("fine-tuned full-loop gate", text)
        self.assertIn("仿真不能证明 fine-tune 必要或闭环成功", text)
        self.assertIn("130 s accumulated active time", text)
        self.assertIn("simulator reset = 0", text)
        self.assertNotIn("comparison_strict_fingerprint_fuzzy_2x2.mp4", text)
        self.assertNotIn("comparison_strict_fingerprint_clear_2x2.mp4", text)
        self.assertNotIn("physical_funnel_guard_ablation.png", text)
        self.assertNotIn("七种方法", text)
        self.assertNotIn("5/12→4/12", text)
        self.assertNotIn("pending", text)

    def test_build_manifest_hashes_and_embedded_video_slides_match(self) -> None:
        manifest = self.manifest
        text = MARKDOWN.read_text(encoding="utf-8")
        self.assertEqual(manifest["slides"], _markdown_slide_count(text))
        self.assertEqual(manifest["missing_local_references"], [])
        self.assertEqual(
            manifest["embedded_video_slides"], [10, 14, 19, 20]
        )
        for label, path in (("markdown", MARKDOWN), ("pptx", PPTX), ("pdf", PDF)):
            with self.subTest(artifact=label):
                record = manifest["artifacts"][label]
                self.assertEqual(Path(record["path"]), path)
                self.assertEqual(record["bytes"], path.stat().st_size)
                self.assertEqual(record["sha256"], _sha256(path))

        expected = {record["slide"]: record for record in manifest["videos"]}
        presentation = Presentation(str(PPTX))
        actual: dict[int, set[str]] = {}
        for slide_number, slide in enumerate(presentation.slides, start=1):
            hashes = {
                hashlib.sha256(rel.target_part.blob).hexdigest()
                for rel in slide.part.rels.values()
                if "video" in rel.reltype
            }
            if hashes:
                actual[slide_number] = hashes
        self.assertEqual(set(actual), set(expected))
        for slide_number, record in expected.items():
            with self.subTest(slide=slide_number):
                source = Path(record["path"])
                self.assertEqual(record["sha256"], _sha256(source))
                self.assertEqual(actual[slide_number], {record["sha256"]})


if __name__ == "__main__":
    unittest.main()
