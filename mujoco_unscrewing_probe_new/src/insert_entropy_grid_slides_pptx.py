#!/usr/bin/env python3
"""Insert the M8 entropy-grid map and embedded explainer into the active deck."""

from __future__ import annotations

import argparse
from io import BytesIO
import os
from pathlib import Path
import posixpath
import tempfile
import xml.etree.ElementTree as ET
import zipfile

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches


CANVAS = (2560, 1440)
FONT_REGULAR = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")
FONT_BOLD = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc")
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--pptx", type=Path, required=True)
    parser.add_argument("--markdown", type=Path, required=True)
    parser.add_argument("--figure", type=Path, required=True)
    parser.add_argument("--video", type=Path, required=True)
    parser.add_argument("--poster", type=Path, required=True)
    return parser.parse_args()


def font(path: Path, size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(path), size)


def _fit(image: Image.Image, box: tuple[int, int, int, int]) -> tuple[Image.Image, tuple[int, int]]:
    left, top, right, bottom = box
    maximum = (right - left, bottom - top)
    copy = image.copy()
    copy.thumbnail(maximum, Image.Resampling.LANCZOS)
    return copy, (left + (maximum[0] - copy.width) // 2, top + (maximum[1] - copy.height) // 2)


def render_map_slide(figure: Path) -> Image.Image:
    source = Image.open(figure).convert("RGB")
    result = source.resize(CANVAS, Image.Resampling.LANCZOS)
    draw = ImageDraw.Draw(result)
    draw.rectangle((2390, 1340, 2560, 1440), fill="white")
    draw.text((2520, 1394), "24", font=font(FONT_REGULAR, 42), fill=(112, 112, 112), anchor="rs")
    return result


def render_video_slide(poster: Path) -> tuple[Image.Image, tuple[int, int, int, int]]:
    image = Image.new("RGB", CANVAS, "white")
    draw = ImageDraw.Draw(image)
    title = font(FONT_BOLD, 61)
    section = font(FONT_BOLD, 34)
    body = font(FONT_REGULAR, 27)
    note = font(FONT_REGULAR, 25)
    draw.text((88, 105), "熵网格提供在线重排覆盖的动作原语", font=title, fill=(23, 55, 94))

    video_box = (88, 240, 1710, 1152)
    source = Image.open(poster).convert("RGB")
    fitted, position = _fit(source, video_box)
    image.paste(fitted, position)
    actual_box = (position[0], position[1], position[0] + fitted.width, position[1] + fitted.height)
    draw.rectangle(actual_box, outline=(51, 65, 85), width=3)
    draw.rounded_rectangle((1770, 240, 2470, 625), radius=12, fill=(248, 250, 252), outline=(203, 213, 225), width=2)
    draw.text((1810, 275), "Pure entropy-grid", font=section, fill=(91, 33, 182))
    lines = (
        "• 从宽视觉先验开始",
        "• 最大化接触信息增益",
        "• 是强 active-perception diagnostic",
        "• 不进入主 baseline ladder",
    )
    y = 355
    for line in lines:
        draw.text((1810, y), line, font=body, fill=(17, 24, 39))
        y += 62

    draw.rounded_rectangle((1770, 675, 2470, 1125), radius=12, fill=(240, 253, 250), outline=(22, 132, 123), width=3)
    draw.text((1810, 710), "Proposed+", font=section, fill=(15, 118, 110))
    lines = (
        "• 继承 conical posterior shortcut",
        "• 主目标：segment capture probability",
        "• information gain 只消除歧义",
        "• miss/contact 后重排覆盖",
        "• capture / unsafe / budget 明确退出",
    )
    y = 790
    for line in lines:
        draw.text((1810, y), line, font=body, fill=(17, 24, 39))
        y += 62

    draw.rectangle((88, 1215, 2470, 1332), fill=(238, 242, 255))
    draw.rectangle((88, 1215, 98, 1332), fill=(37, 99, 235))
    draw.text(
        (125, 1248),
        "合理闭环不要求每次成功：观测改变 belief，belief 改变下一动作，失败可归因且安全退出。",
        font=note,
        fill=(17, 24, 39),
    )
    draw.text((2520, 1394), "25", font=font(FONT_REGULAR, 42), fill=(112, 112, 112), anchor="rs")
    return image, actual_box


def image_stream(image: Image.Image) -> BytesIO:
    stream = BytesIO()
    image.save(stream, format="PNG", optimize=True)
    stream.seek(0)
    return stream


def primary_raster_path(archive: zipfile.ZipFile, slide_number: int) -> str:
    slide_path = f"ppt/slides/slide{slide_number}.xml"
    rels_path = f"ppt/slides/_rels/slide{slide_number}.xml.rels"
    slide_root = ET.fromstring(archive.read(slide_path))
    relationships = ET.fromstring(archive.read(rels_path))
    blip = slide_root.find(f".//{{{P_NS}}}bg//{{{A_NS}}}blip")
    if blip is None:
        blip = slide_root.find(f".//{{{P_NS}}}spTree/{{{P_NS}}}pic//{{{A_NS}}}blip")
    if blip is None:
        raise RuntimeError(f"slide {slide_number} has no primary raster")
    relation_id = blip.attrib[f"{{{R_NS}}}embed"]
    for relationship in relationships.findall(f"{{{REL_NS}}}Relationship"):
        if relationship.attrib.get("Id") == relation_id:
            return posixpath.normpath(posixpath.join("ppt/slides", relationship.attrib["Target"]))
    raise RuntimeError(f"slide {slide_number} lost raster relationship {relation_id}")


def renumber(blob: bytes, page: int) -> bytes:
    image = Image.open(BytesIO(blob)).convert("RGB")
    draw = ImageDraw.Draw(image)
    width, height = image.size
    draw.rectangle((width - 220, height - 130, width, height), fill="white")
    draw.text(
        (width - 40, height - 46),
        str(page),
        font=font(FONT_REGULAR, max(28, int(round(height * 42 / 1440)))),
        fill=(112, 112, 112),
        anchor="rs",
    )
    stream = BytesIO()
    image.save(stream, format="PNG", optimize=True)
    return stream.getvalue()


def replace_zip_entries(path: Path, replacements: dict[str, bytes]) -> None:
    with tempfile.NamedTemporaryFile(prefix=f".{path.stem}.entropy.", suffix=".pptx", dir=path.parent, delete=False) as handle:
        temporary_path = Path(handle.name)
    try:
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary_path, "w") as target:
            missing = set(replacements) - set(source.namelist())
            if missing:
                raise RuntimeError(f"saved PPTX lost raster media: {sorted(missing)}")
            for info in source.infolist():
                target.writestr(info, replacements.get(info.filename, source.read(info.filename)))
        os.replace(temporary_path, path)
    finally:
        temporary_path.unlink(missing_ok=True)


def main() -> None:
    args = parse_args()
    for path in (args.pptx, args.markdown, args.figure, args.video, args.poster, FONT_REGULAR, FONT_BOLD):
        if not path.is_file():
            raise FileNotFoundError(path)
    chunks = args.markdown.read_text(encoding="utf-8").split("\n---\n")[1:]
    if len(chunks) != 49:
        raise ValueError(f"expected 49-slide updated source, got {len(chunks)}")

    presentation = Presentation(str(args.pptx))
    if len(presentation.slides) != 47:
        raise ValueError(f"expected 47-slide active deck, got {len(presentation.slides)}")

    replacements: dict[str, bytes] = {}
    with zipfile.ZipFile(args.pptx, "r") as archive:
        for old_page in range(24, 48):
            new_page = old_page + 2
            if "<!-- _paginate: false -->" in chunks[new_page - 1]:
                continue
            media_path = primary_raster_path(archive, old_page)
            replacements[media_path] = renumber(archive.read(media_path), new_page)

    blank = presentation.slide_layouts[0]
    map_slide = presentation.slides.add_slide(blank)
    for shape in list(map_slide.shapes):
        shape._element.getparent().remove(shape._element)
    map_slide.shapes.add_picture(image_stream(render_map_slide(args.figure)), 0, 0, width=presentation.slide_width, height=presentation.slide_height)

    video_slide = presentation.slides.add_slide(blank)
    for shape in list(video_slide.shapes):
        shape._element.getparent().remove(shape._element)
    background, video_box = render_video_slide(args.poster)
    video_slide.shapes.add_picture(image_stream(background), 0, 0, width=presentation.slide_width, height=presentation.slide_height)
    left, top, right, bottom = video_box
    video_slide.shapes.add_movie(
        str(args.video),
        Inches(13.333 * left / CANVAS[0]),
        Inches(7.5 * top / CANVAS[1]),
        Inches(13.333 * (right - left) / CANVAS[0]),
        Inches(7.5 * (bottom - top) / CANVAS[1]),
        poster_frame_image=str(args.poster),
        mime_type="video/mp4",
    )

    slide_ids = presentation.slides._sldIdLst
    map_id = slide_ids[-2]
    video_id = slide_ids[-1]
    slide_ids.remove(video_id)
    slide_ids.remove(map_id)
    slide_ids.insert(23, map_id)
    slide_ids.insert(24, video_id)

    with tempfile.NamedTemporaryFile(prefix=f".{args.pptx.stem}.", suffix=".pptx", dir=args.pptx.parent, delete=False) as handle:
        temporary_path = Path(handle.name)
    try:
        presentation.save(str(temporary_path))
        os.replace(temporary_path, args.pptx)
    finally:
        temporary_path.unlink(missing_ok=True)
    replace_zip_entries(args.pptx, replacements)
    print("inserted entropy-grid slides 24/25 and embedded MP4; deck 47 -> 49 slides")


if __name__ == "__main__":
    main()
