#!/usr/bin/env python3
"""Deterministically build the active Marp deck, embed videos, and export PDF."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile

from PIL import Image
from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_MARKDOWN = (
    ROOT.parent / "group_meeting_trace/unscrewing_contact_state_pipeline_slides_0709.md"
)
# Bounding regions occupied by the linked poster in the Marp-rendered slide.
# The actual movie box is aspect-fitted inside these regions below.
VIDEO_REGION_WIDE = (2.72, 1.34, 7.90, 4.45)
VIDEO_REGION_TWO_COLUMN = (0.44, 1.34, 6.14, 4.61)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--markdown", type=Path, default=DEFAULT_MARKDOWN)
    parser.add_argument("--pptx", type=Path)
    parser.add_argument("--pdf", type=Path)
    return parser.parse_args()


def _latest(paths: list[Path], label: str) -> Path:
    candidates = [path for path in paths if path.is_file()]
    if not candidates:
        raise FileNotFoundError(f"cannot locate {label}")
    return max(candidates, key=lambda path: path.stat().st_mtime_ns)


def _toolchain() -> tuple[Path, Path, Path]:
    node = _latest(
        list(Path("/home/njy/.vscode-server/cli/servers").glob("Stable-*/server/node")),
        "VS Code server Node",
    )
    marp = _latest(
        list(
            Path("/home/njy/.vscode-server/extensions").glob(
                "marp-team.marp-vscode-*/node_modules/@marp-team/marp-cli/marp-cli.js"
            )
        ),
        "Marp CLI",
    )
    chrome = _latest(
        list(
            Path("/home/njy/.cache/puppeteer-browsers/chrome-headless-shell").glob(
                "*/chrome-headless-shell-linux64/chrome-headless-shell"
            )
        ),
        "Chrome headless shell",
    )
    return node, marp, chrome


def _chunks(markdown: Path) -> list[str]:
    parts = markdown.read_text(encoding="utf-8").split("\n---\n")
    if len(parts) < 2:
        raise ValueError("Markdown does not contain Marp slide separators")
    return parts[1:]


def _local_path(markdown: Path, reference: str) -> Path:
    return (markdown.parent / reference).resolve()


def _local_references(markdown: Path) -> list[Path]:
    text = markdown.read_text(encoding="utf-8")
    references = re.findall(r'(?:src|href)="([^"]+)"', text)
    references.extend(re.findall(r"\]\(([^)]+)\)", text))
    paths: list[Path] = []
    for reference in references:
        clean = reference.strip().split("#", 1)[0]
        if not clean or re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*:", clean):
            continue
        path = _local_path(markdown, clean)
        if path not in paths:
            paths.append(path)
    return paths


def _fit_box(region: tuple[float, float, float, float], aspect: float) -> tuple[float, float, float, float]:
    """Aspect-fit a movie inside the poster region without covering nearby text."""

    left, top, width, height = region
    if width / height > aspect:
        fitted_width = height * aspect
        return (left + (width - fitted_width) / 2.0, top, fitted_width, height)
    fitted_height = width / aspect
    return (left, top + (height - fitted_height) / 2.0, width, fitted_height)


def _video_specs(markdown: Path, chunks: list[str]) -> list[dict]:
    specs: list[dict] = []
    for slide_number, chunk in enumerate(chunks, start=1):
        videos = list(dict.fromkeys(re.findall(r'href="([^"]+\.mp4)"', chunk)))
        if not videos:
            continue
        if len(videos) != 1:
            raise ValueError(f"slide {slide_number} has {len(videos)} distinct videos")
        images = re.findall(r'<img[^>]+src="([^"]+\.(?:png|jpg|jpeg))"', chunk)
        if not images:
            raise ValueError(f"slide {slide_number} has a video but no poster image")
        video = _local_path(markdown, videos[0])
        poster = _local_path(markdown, images[0])
        if not video.is_file() or not poster.is_file():
            raise FileNotFoundError((video, poster))
        video_position = chunk.find(videos[0])
        columns_position = chunk.find('class="cols"')
        # Use the two-column box only when the movie itself is inside a columns
        # block.  Explanatory columns that appear *after* a full-width movie do
        # not change its placement.
        region = (
            VIDEO_REGION_TWO_COLUMN
            if 0 <= columns_position < video_position
            else VIDEO_REGION_WIDE
        )
        with Image.open(poster) as poster_image:
            aspect = poster_image.width / poster_image.height
        box = _fit_box(region, aspect)
        specs.append(
            {
                "slide": slide_number,
                "video": video,
                "poster": poster,
                "box_inches": box,
            }
        )
    return specs


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _embed_videos(pptx: Path, specs: list[dict]) -> None:
    presentation = Presentation(str(pptx))
    for spec in specs:
        slide = presentation.slides[int(spec["slide"]) - 1]
        left, top, width, height = spec["box_inches"]
        movie = slide.shapes.add_movie(
            str(spec["video"]),
            Inches(left),
            Inches(top),
            Inches(width),
            Inches(height),
            poster_frame_image=str(spec["poster"]),
            mime_type="video/mp4",
        )
        movie.name = f"Embedded mechanism video slide {spec['slide']}"
    with tempfile.NamedTemporaryFile(
        prefix=f".{pptx.stem}.", suffix=".pptx", dir=pptx.parent, delete=False
    ) as handle:
        temporary = Path(handle.name)
    try:
        presentation.save(str(temporary))
        os.replace(temporary, pptx)
    finally:
        temporary.unlink(missing_ok=True)


def _pdf_pages(path: Path) -> int:
    output = subprocess.check_output(["pdfinfo", str(path)], text=True)
    match = re.search(r"^Pages:\s+(\d+)$", output, flags=re.MULTILINE)
    if match is None:
        raise RuntimeError("pdfinfo did not report a page count")
    return int(match.group(1))


def main() -> None:
    args = parse_args()
    markdown = args.markdown.resolve()
    pptx = (args.pptx or markdown.with_suffix(".pptx")).resolve()
    pdf = (args.pdf or markdown.with_suffix(".pdf")).resolve()
    if not markdown.is_file():
        raise FileNotFoundError(markdown)
    chunks = _chunks(markdown)
    local_references = _local_references(markdown)
    missing_references = [path for path in local_references if not path.is_file()]
    if missing_references:
        raise FileNotFoundError(
            "missing local deck references: "
            + ", ".join(str(path) for path in missing_references)
        )
    video_specs = _video_specs(markdown, chunks)
    node, marp, chrome = _toolchain()
    pptx.parent.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory(prefix="m8_deck_build_") as temporary_dir:
        temporary = Path(temporary_dir)
        raw_pptx = temporary / pptx.name
        subprocess.run(
            [
                str(node),
                str(marp),
                "--allow-local-files",
                "--browser",
                "chrome",
                "--browser-path",
                str(chrome),
                "--pptx",
                str(markdown),
                "-o",
                str(raw_pptx),
            ],
            cwd=markdown.parent,
            check=True,
        )
        if len(Presentation(str(raw_pptx)).slides) != len(chunks):
            raise RuntimeError("Marp slide count does not match Markdown")
        # Export the poster-only Marp package before embedding H.264 movies.
        # LibreOffice tries to initialize a system video decoder for embedded
        # movies during headless PDF export; on machines without that codec it
        # can terminate with SIGSEGV even though the slide poster is complete.
        # The raw package already contains the exact poster visible in the PDF.
        pdf_out = temporary / "pdf"
        pdf_out.mkdir()
        lo_profile = temporary / "lo-profile"
        subprocess.run(
            [
                "soffice",
                f"-env:UserInstallation=file://{lo_profile}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(pdf_out),
                str(raw_pptx),
            ],
            check=True,
        )
        converted = pdf_out / raw_pptx.with_suffix(".pdf").name
        if not converted.is_file():
            raise RuntimeError("LibreOffice did not create the PDF")
        shutil.copy2(converted, pdf)
        pdf.chmod(0o644)

        shutil.copy2(raw_pptx, pptx)
        _embed_videos(pptx, video_specs)
        # NamedTemporaryFile creates the replacement package as 0600.  Restore
        # ordinary read permissions so the built deck is shareable without a
        # second manual chmod step.
        pptx.chmod(0o644)

    presentation = Presentation(str(pptx))
    embedded_slides = []
    for index, slide in enumerate(presentation.slides, start=1):
        if any("video" in rel.reltype for rel in slide.part.rels.values()):
            embedded_slides.append(index)
    if len(presentation.slides) != len(chunks) or _pdf_pages(pdf) != len(chunks):
        raise RuntimeError("final deck page counts are inconsistent")
    expected_slides = [int(spec["slide"]) for spec in video_specs]
    if embedded_slides != expected_slides:
        raise RuntimeError(
            f"embedded video slides {embedded_slides}, expected {expected_slides}"
        )

    manifest = {
        "markdown": str(markdown),
        "pptx": str(pptx),
        "pdf": str(pdf),
        "artifacts": {
            label: {
                "path": str(path),
                "bytes": path.stat().st_size,
                "sha256": _sha256(path),
            }
            for label, path in {
                "markdown": markdown,
                "pptx": pptx,
                "pdf": pdf,
            }.items()
        },
        "slides": len(chunks),
        "local_reference_count": len(local_references),
        "missing_local_references": [],
        "embedded_video_slides": embedded_slides,
        "videos": [
            {
                "slide": int(spec["slide"]),
                "path": str(spec["video"]),
                "sha256": _sha256(spec["video"]),
            }
            for spec in video_specs
        ],
        "toolchain": {
            "node": str(node),
            "marp": str(marp),
            "chrome": str(chrome),
            "libreoffice": shutil.which("soffice"),
        },
    }
    manifest_path = markdown.with_name(f"{markdown.stem}_build_manifest.json")
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
