#!/usr/bin/env python3
"""Compact the physical-funnel section of the active raster-based Marp PPTX."""

from __future__ import annotations

import argparse
from io import BytesIO
import os
from pathlib import Path
import posixpath
import tempfile
import xml.etree.ElementTree as ET
import zipfile

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation


CANVAS = (2560, 1440)
FONT_REGULAR = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")
FONT_BOLD = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc")

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--pptx", type=Path, required=True)
    parser.add_argument("--markdown", type=Path, required=True)
    return parser.parse_args()


def font(path: Path, size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(path), size)


def draw_wrapped(
    draw: ImageDraw.ImageDraw,
    text: str,
    xy: tuple[int, int],
    text_font: ImageFont.FreeTypeFont,
    width: int,
    fill: tuple[int, int, int] = (17, 24, 39),
    spacing: int = 10,
) -> int:
    lines: list[str] = []
    current = ""
    for character in text:
        candidate = current + character
        if current and draw.textbbox((0, 0), candidate, font=text_font)[2] > width:
            lines.append(current)
            current = character
        else:
            current = candidate
    if current:
        lines.append(current)
    x, y = xy
    line_height = draw.textbbox((0, 0), "国Ag", font=text_font)[3] + spacing
    for line in lines:
        draw.text((x, y), line, font=text_font, fill=fill)
        y += line_height
    return y


def render_merged_control_slide() -> Image.Image:
    image = Image.new("RGB", CANVAS, "white")
    draw = ImageDraw.Draw(image)
    title_font = font(FONT_BOLD, 63)
    section_font = font(FONT_BOLD, 40)
    body_font = font(FONT_REGULAR, 31)
    formula_font = font(FONT_REGULAR, 36)
    chain_font = font(FONT_BOLD, 33)
    note_font = font(FONT_REGULAR, 29)
    page_font = font(FONT_REGULAR, 42)

    draw.text(
        (88, 135),
        "捕获后的控制骨架：tip-RCC + bounded recovery",
        font=title_font,
        fill=(23, 55, 94),
    )

    left = (88, 255, 1248, 1090)
    right = (1312, 255, 2472, 1090)
    for box in (left, right):
        draw.rounded_rectangle(box, radius=10, fill=(248, 250, 252), outline=(203, 213, 225), width=2)

    draw.rectangle((88, 255, 1248, 335), fill=(234, 241, 248))
    draw.text((122, 272), "1  顺应执行：让环境约束帮助回正", font=section_font, fill=(23, 55, 94))
    draw.text((132, 385), "M_a e'' + D_a e' + K_a e = w_ext - w_d", font=formula_font, fill=(17, 24, 39))
    draw.text((132, 470), "tau_tip = tau_s + (p_s - p_tip) x F_s", font=formula_font, fill=(17, 24, 39))
    y = 585
    for line in (
        "• capture cue 成立后冻结 XY 搜索",
        "• command upright，并维持小轴向力",
        "• 降低平面刚度，把顺应中心移到 peg tip",
        "• 倒角/内壁法向反力负责被动找正",
    ):
        draw.text((132, y), line, font=body_font, fill=(17, 24, 39))
        y += 88

    draw.rectangle((1312, 255, 2472, 335), fill=(236, 253, 245))
    draw.text((1346, 272), "2  独立监督：只处理不可恢复趋势", font=section_font, fill=(15, 118, 110))
    draw.text((1356, 385), "持续监测：", font=chain_font, fill=(17, 24, 39))
    y = 460
    for line in (
        "• depth progress 与 CAD terminal band",
        "• actual tilt、F/T wrench 与接触功 E_c",
        "• 无进展 + 越界，而不是单帧峰值",
    ):
        draw.text((1356, y), line, font=body_font, fill=(17, 24, 39))
        y += 88
    draw.rounded_rectangle((1350, 755, 2435, 960), radius=10, fill=(240, 253, 250), outline=(22, 132, 123), width=3)
    draw.text((1385, 785), "最多一次 bounded recovery", font=chain_font, fill=(15, 118, 110))
    draw_wrapped(
        draw,
        "卸载 → 回撤 → 回正 → 低力重触 → 继续或 safe abort",
        (1385, 850),
        chain_font,
        990,
        fill=(17, 24, 39),
        spacing=8,
    )

    draw.rectangle((88, 1160, 2472, 1305), fill=(255, 251, 235))
    draw.rectangle((88, 1160, 98, 1305), fill=(217, 119, 6))
    draw_wrapped(
        draw,
        "实现边界：MuJoCo 是 Cartesian tip-RCC/supervisor proxy；真机由 iiwa impedance + 外置 F/T 承载。仿真 gain、阈值和恢复成功率不能直接迁移。",
        (120, 1191),
        note_font,
        2260,
        spacing=8,
    )
    draw.text((2520, 1394), "30", font=page_font, fill=(112, 112, 112), anchor="rs")
    return image


def slide_chunks(markdown_path: Path) -> list[str]:
    return markdown_path.read_text(encoding="utf-8").split("\n---\n")[1:]


def primary_raster_path(archive: zipfile.ZipFile, slide_number: int) -> str:
    slide_path = f"ppt/slides/slide{slide_number}.xml"
    rels_path = f"ppt/slides/_rels/slide{slide_number}.xml.rels"
    slide_root = ET.fromstring(archive.read(slide_path))
    relationships = ET.fromstring(archive.read(rels_path))

    blip = slide_root.find(f".//{{{P_NS}}}bg//{{{A_NS}}}blip")
    if blip is None:
        blip = slide_root.find(f".//{{{P_NS}}}spTree/{{{P_NS}}}pic//{{{A_NS}}}blip")
    if blip is None:
        raise RuntimeError(f"slide {slide_number} has no primary raster")
    relation_id = blip.attrib[f"{{{R_NS}}}embed"]

    for relationship in relationships.findall(f"{{{REL_NS}}}Relationship"):
        if relationship.attrib.get("Id") == relation_id:
            return posixpath.normpath(
                posixpath.join("ppt/slides", relationship.attrib["Target"])
            )
    raise RuntimeError(f"slide {slide_number} lost raster relationship {relation_id}")


def renumber(image_blob: bytes, page: int) -> bytes:
    image = Image.open(BytesIO(image_blob)).convert("RGB")
    draw = ImageDraw.Draw(image)
    draw.rectangle((2340, 1310, 2560, 1440), fill="white")
    draw.text(
        (2520, 1394),
        str(page),
        font=font(FONT_REGULAR, 42),
        fill=(112, 112, 112),
        anchor="rs",
    )
    stream = BytesIO()
    image.save(stream, format="PNG", optimize=True)
    return stream.getvalue()


def delete_slide(presentation: Presentation, one_based_index: int) -> None:
    slide_ids = presentation.slides._sldIdLst
    slide_id = slide_ids[one_based_index - 1]
    presentation.part.drop_rel(slide_id.rId)
    slide_ids.remove(slide_id)


def replace_zip_entries(path: Path, replacements: dict[str, bytes]) -> None:
    with tempfile.NamedTemporaryFile(
        prefix=f".{path.stem}.compact.", suffix=".pptx", dir=path.parent, delete=False
    ) as handle:
        temporary_path = Path(handle.name)
    try:
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary_path, "w") as target:
            missing = set(replacements) - set(source.namelist())
            if missing:
                raise RuntimeError(f"saved PPTX lost raster media: {sorted(missing)}")
            for info in source.infolist():
                target.writestr(info, replacements.get(info.filename, source.read(info.filename)))
        os.replace(temporary_path, path)
    finally:
        temporary_path.unlink(missing_ok=True)


def main() -> None:
    args = parse_args()
    for path in (args.pptx, args.markdown, FONT_REGULAR, FONT_BOLD):
        if not path.is_file():
            raise FileNotFoundError(path)
    chunks = slide_chunks(args.markdown)
    if len(chunks) != 47:
        raise ValueError(f"expected compacted 47-slide source, got {len(chunks)}")

    presentation = Presentation(str(args.pptx))
    if len(presentation.slides) != 49:
        raise ValueError(f"expected pre-compaction 49-slide deck, got {len(presentation.slides)}")

    replacements: dict[str, bytes] = {}
    with zipfile.ZipFile(args.pptx, "r") as archive:
        for old_page in range(33, 50):
            new_page = old_page - 2
            if "<!-- _paginate: false -->" in chunks[new_page - 1]:
                continue
            media_path = primary_raster_path(archive, old_page)
            replacements[media_path] = renumber(archive.read(media_path), new_page)

    merged_slide = presentation.slides[30]
    for shape in list(merged_slide.shapes):
        shape._element.getparent().remove(shape._element)
    merged = render_merged_control_slide()
    stream = BytesIO()
    merged.save(stream, format="PNG", optimize=True)
    stream.seek(0)
    merged_slide.shapes.add_picture(
        stream,
        0,
        0,
        width=presentation.slide_width,
        height=presentation.slide_height,
    )

    # Delete the former standalone shield page first, then the duplicated cue
    # page. The merged former control page shifts from 31 to 30.
    delete_slide(presentation, 32)
    delete_slide(presentation, 30)

    with tempfile.NamedTemporaryFile(
        prefix=f".{args.pptx.stem}.", suffix=".pptx", dir=args.pptx.parent, delete=False
    ) as handle:
        temporary_path = Path(handle.name)
    try:
        presentation.save(str(temporary_path))
        os.replace(temporary_path, args.pptx)
    finally:
        temporary_path.unlink(missing_ok=True)

    replace_zip_entries(args.pptx, replacements)
    print("compacted physical-funnel section: 49 -> 47 slides")


if __name__ == "__main__":
    main()
