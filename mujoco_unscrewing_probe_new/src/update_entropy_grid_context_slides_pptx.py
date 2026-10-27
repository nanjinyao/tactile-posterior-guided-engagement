#!/usr/bin/env python3
"""Synchronize finite-budget coverage and Proposed+ framing in the active PPTX."""

from __future__ import annotations

import argparse
from io import BytesIO
import os
from pathlib import Path
import tempfile

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation


CANVAS = (2560, 1440)
FONT_REGULAR = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")
FONT_BOLD = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--pptx", type=Path, required=True)
    parser.add_argument("--benchmark", type=Path, required=True)
    return parser.parse_args()


def font(path: Path, size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(path), size)


def base(title: str, page: int) -> tuple[Image.Image, ImageDraw.ImageDraw]:
    image = Image.new("RGB", CANVAS, "white")
    draw = ImageDraw.Draw(image)
    draw.text((88, 105), title, font=font(FONT_BOLD, 61), fill=(23, 55, 94))
    draw.text((2520, 1394), str(page), font=font(FONT_REGULAR, 42), fill=(112, 112, 112), anchor="rs")
    return image, draw


def render_coverage() -> Image.Image:
    image, draw = base("有限域、有限预算内的覆盖来自实体宽度", 23)
    section = font(FONT_BOLD, 37)
    formula = font(FONT_REGULAR, 48)
    body = font(FONT_REGULAR, 31)
    note = font(FONT_REGULAR, 27)
    draw.rounded_rectangle((88, 245, 1215, 900), radius=12, fill=(248, 250, 252), outline=(203, 213, 225), width=2)
    draw.text((130, 285), "1  螺旋的空间条件", font=section, fill=(23, 55, 94))
    draw.text((165, 405), "r(θ) = bθ", font=formula, fill=(17, 24, 39))
    draw.text((165, 515), "Δr_turn = 2πb ≤ 2r_eff", font=formula, fill=(17, 24, 39))
    draw.text((135, 675), "轨迹的 Minkowski tube 不留下", font=body, fill=(17, 24, 39))
    draw.text((135, 735), "大于捕获直径的径向空档。", font=body, fill=(17, 24, 39))

    draw.rounded_rectangle((1300, 245, 2470, 900), radius=12, fill=(240, 253, 250), outline=(22, 132, 123), width=3)
    draw.text((1342, 285), "2  有限时间的完整条件", font=section, fill=(15, 118, 110))
    draw.text((1365, 405), "R covered by Union B(x, r_eff)", font=font(FONT_REGULAR, 39), fill=(17, 24, 39))
    draw.text((1545, 465), "x in Gamma_T", font=body, fill=(17, 24, 39))
    draw.text((1365, 575), "T >= path_length(Gamma_T)/v", font=font(FONT_REGULAR, 39), fill=(17, 24, 39))
    draw.text((1545, 635), "+ T_dwell", font=font(FONT_REGULAR, 39), fill=(17, 24, 39))
    draw.text((1348, 725), "还要求 capture cue 可用、轨迹不中断，", font=body, fill=(17, 24, 39))
    draw.text((1348, 785), "并且剩余 active budget 足够。", font=body, fill=(17, 24, 39))

    draw.rectangle((88, 980, 2470, 1185), fill=(238, 242, 255))
    draw.rectangle((88, 980, 98, 1185), fill=(37, 99, 235))
    draw.text((130, 1018), "Learning 只缩小 residual search area；低置信时运输趋零，", font=body, fill=(17, 24, 39))
    draw.text((130, 1080), "并在剩余预算允许时退化为 full outward coverage。", font=body, fill=(17, 24, 39))
    draw.rectangle((88, 1230, 2470, 1335), fill=(255, 251, 235))
    draw.text((122, 1260), "条件覆盖只保证轨迹管与 capture region 相交，不等于每次完成 CAD-bottom；主审计 T_max = 130 s。", font=note, fill=(120, 65, 0))
    return image


def _table_cell(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], text: str, *, header: bool = False, accent: bool = False) -> None:
    fill = (234, 241, 248) if header else ((240, 253, 250) if accent else (255, 255, 255))
    draw.rectangle(box, fill=fill, outline=(203, 213, 225), width=2)
    selected_font = font(FONT_BOLD if header or accent else FONT_REGULAR, 27 if not header else 29)
    left, top, right, bottom = box
    lines = text.split("\n")
    y = top + 17
    for line in lines:
        draw.text((left + 16, y), line, font=selected_font, fill=(15, 118, 110) if accent else (23, 55, 94) if header else (17, 24, 39))
        y += 42


def render_ladder() -> Image.Image:
    image, draw = base("方法演化与主 baseline：每一条只改变一个机制", 27)
    x = [88, 430, 955, 1720, 2470]
    y = [245, 335, 510, 685, 860, 1035]
    headers = ("方法", "感知", "搜索", "研究问题")
    rows = (
        ("Proposed", "一次 learned fingerprint", "bounded move +\nresidual/full outward", "方向捷径能否保留覆盖"),
        ("Proposed+", "相同 fingerprint\n+ miss/contact", "task-aware adaptive\ncoverage + full outward", "belief 更新能否减少\n无效覆盖与接触暴露"),
        ("A outward", "无 probe", "视觉锚点\noutward spiral", "不使用主动触觉的成本"),
        ("B direct/retry", "相同 fingerprint", "direct ray；失败后\n偏移重探，最多 3 次", "spiral fallback 是否必要"),
    )
    for column, label in enumerate(headers):
        _table_cell(draw, (x[column], y[0], x[column + 1], y[1]), label, header=True)
    for row_index, row in enumerate(rows, start=1):
        for column, value in enumerate(row):
            _table_cell(draw, (x[column], y[row_index], x[column + 1], y[row_index + 1]), value, accent=row_index == 2)
    draw.rectangle((88, 1125, 2470, 1325), fill=(248, 250, 252), outline=(203, 213, 225), width=2)
    body = font(FONT_REGULAR, 28)
    draw.text((125, 1162), "统一：相同初态 / nuisance、capture / terminal、安全层与 130 s active budget。", font=body, fill=(17, 24, 39))
    draw.text((125, 1225), "Pure entropy-grid 与 random 保留为附录诊断；Oracle 只作上界。", font=body, fill=(17, 24, 39))
    return image


def render_benchmark(benchmark: Path) -> Image.Image:
    image, draw = base("M8 冻结 36 点：优势是覆盖、时间与接触成本的平衡", 28)
    source = Image.open(benchmark).convert("RGB")
    source.thumbnail((2180, 810), Image.Resampling.LANCZOS)
    x = (CANVAS[0] - source.width) // 2
    image.paste(source, (x, 220))
    draw.rounded_rectangle((88, 1065, 2470, 1325), radius=8, fill=(240, 253, 250), outline=(22, 132, 123), width=3)
    body = font(FONT_REGULAR, 29)
    bold = font(FONT_BOLD, 30)
    draw.text((125, 1102), "T_max = 130 s active budget：Proposed 34/36；A 30/36；B 28/36；random 19/36。", font=bold, fill=(15, 118, 110))
    draw.text((125, 1170), "正确 claim：一次 learned fingerprint 在覆盖、速度与接触成本之间取得更好平衡，", font=body, fill=(17, 24, 39))
    draw.text((125, 1228), "不是每个共同成功点都最快；未完成需分解 timeout / safe abort / terminal failure。", font=body, fill=(17, 24, 39))
    return image


def replace_slide(slide, image: Image.Image, width: int, height: int) -> None:
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    stream = BytesIO()
    image.save(stream, format="PNG", optimize=True)
    stream.seek(0)
    slide.shapes.add_picture(stream, 0, 0, width=width, height=height)


def main() -> None:
    args = parse_args()
    for path in (args.pptx, args.benchmark, FONT_REGULAR, FONT_BOLD):
        if not path.is_file():
            raise FileNotFoundError(path)
    presentation = Presentation(str(args.pptx))
    if len(presentation.slides) != 49:
        raise ValueError(f"expected 49 slides, got {len(presentation.slides)}")
    replacements = {23: render_coverage(), 27: render_ladder(), 28: render_benchmark(args.benchmark)}
    for page, image in replacements.items():
        replace_slide(presentation.slides[page - 1], image, presentation.slide_width, presentation.slide_height)
    with tempfile.NamedTemporaryFile(prefix=f".{args.pptx.stem}.context.", suffix=".pptx", dir=args.pptx.parent, delete=False) as handle:
        temporary = Path(handle.name)
    try:
        presentation.save(str(temporary))
        os.replace(temporary, args.pptx)
    finally:
        temporary.unlink(missing_ok=True)
    print("updated PPTX slides 23, 27, and 28")


if __name__ == "__main__":
    main()
