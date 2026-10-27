#!/usr/bin/env python3
"""Embed the active funnel-guard comparison into the exported Marp deck."""

from __future__ import annotations

import argparse
import os
from pathlib import Path
import tempfile

from pptx import Presentation
from pptx.util import Inches


MOVIE_SHAPE_NAME = "Funnel Guard MuJoCo Comparison"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--pptx", type=Path, required=True)
    parser.add_argument("--video", type=Path, required=True)
    parser.add_argument("--poster", type=Path, required=True)
    parser.add_argument("--slide", type=int, default=34, help="One-based slide index")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    for path in (args.pptx, args.video, args.poster):
        if not path.is_file():
            raise FileNotFoundError(path)

    presentation = Presentation(str(args.pptx))
    if not 1 <= args.slide <= len(presentation.slides):
        raise ValueError(f"slide {args.slide} outside 1..{len(presentation.slides)}")

    slide = presentation.slides[args.slide - 1]
    for shape in list(slide.shapes):
        if shape.name == MOVIE_SHAPE_NAME:
            shape._element.getparent().remove(shape._element)

    # Match the 16:9 keyframe box in the 13.333 x 7.5 inch Marp slide.
    movie = slide.shapes.add_movie(
        str(args.video),
        Inches(2.72),
        Inches(1.34),
        Inches(7.90),
        Inches(4.45),
        poster_frame_image=str(args.poster),
        mime_type="video/mp4",
    )
    movie.name = MOVIE_SHAPE_NAME

    with tempfile.NamedTemporaryFile(
        prefix=f".{args.pptx.stem}.", suffix=".pptx", dir=args.pptx.parent, delete=False
    ) as handle:
        temporary_path = Path(handle.name)

    try:
        presentation.save(str(temporary_path))
        os.replace(temporary_path, args.pptx)
    finally:
        temporary_path.unlink(missing_ok=True)

    print(
        f"embedded {args.video.name} on slide {args.slide} "
        f"of {args.pptx.name}"
    )


if __name__ == "__main__":
    main()
