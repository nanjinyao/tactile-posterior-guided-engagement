#!/usr/bin/env python3
"""Insert the extreme physical-funnel stress slide into the active Marp PPTX."""

from __future__ import annotations

import argparse
from io import BytesIO
import os
from pathlib import Path
import posixpath
import tempfile
import xml.etree.ElementTree as ET
import zipfile

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation


CANVAS = (2560, 1440)
INSERT_SLIDE = 35
FONT_REGULAR = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")
FONT_BOLD = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--pptx", type=Path, required=True)
    parser.add_argument("--markdown", type=Path, required=True)
    parser.add_argument("--poster", type=Path, required=True)
    return parser.parse_args()


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont, width: int) -> list[str]:
    lines: list[str] = []
    current = ""
    for character in text:
        candidate = current + character
        if current and draw.textbbox((0, 0), candidate, font=font)[2] > width:
            lines.append(current)
            current = character
        else:
            current = candidate
    if current:
        lines.append(current)
    return lines


def draw_lines(
    draw: ImageDraw.ImageDraw,
    lines: list[str],
    xy: tuple[int, int],
    font: ImageFont.FreeTypeFont,
    fill: tuple[int, int, int],
    spacing: int,
) -> None:
    x, y = xy
    line_height = draw.textbbox((0, 0), "国Ag", font=font)[3] + spacing
    for line in lines:
        draw.text((x, y), line, font=font, fill=fill)
        y += line_height


def render_slide(poster_path: Path) -> Image.Image:
    image = Image.new("RGB", CANVAS, "white")
    draw = ImageDraw.Draw(image)
    title_font = ImageFont.truetype(str(FONT_BOLD), 63)
    evidence_font = ImageFont.truetype(str(FONT_REGULAR), 35)
    note_font = ImageFont.truetype(str(FONT_REGULAR), 27)
    page_font = ImageFont.truetype(str(FONT_REGULAR), 42)

    draw.text(
        (88, 135),
        "极端压力测试：倾斜楔紧，还是卸载后回正？",
        font=title_font,
        fill=(23, 55, 94),
    )

    poster = Image.open(poster_path).convert("RGB")
    poster = poster.resize((1517, 854), Image.Resampling.LANCZOS)
    image.paste(poster, (522, 257))

    evidence_top = 1140
    draw.rectangle((88, evidence_top, 2474, 1301), fill=(240, 253, 250))
    draw.rectangle((88, evidence_top, 98, 1301), fill=(22, 132, 123))
    evidence = (
        "同一极端残余位姿：8° 倾斜命令 + 0.6 mm 横向 TCP 偏差，M8 深孔几何倾角尺度仅约 1.43°。"
        "无漏斗：实际峰值 9.69°，29.24 N / 128.59 mNm 后楔紧 hard stop；tip-RCC + recovery：峰值 5.99°，"
        "一次卸载/回撤/回正/重触后，以 0.01° 到达 CAD bottom。"
    )
    evidence_lines = wrap_text(draw, evidence, evidence_font, 2310)
    draw_lines(draw, evidence_lines, (118, 1161), evidence_font, (17, 24, 39), 8)

    note = (
        "刻意超过 nominal clearance 的配对 MuJoCo mechanism stress test：用于解释 recovery；"
        "不代表真机故障频率，也不证明严格 passivity、一般恢复率或完成率提升。"
    )
    draw_lines(draw, wrap_text(draw, note, note_font, 2280), (88, 1316), note_font, (51, 65, 85), 5)
    draw.text((2520, 1394), str(INSERT_SLIDE), font=page_font, fill=(112, 112, 112), anchor="rs")
    return image


def slide_chunks(markdown_path: Path) -> list[str]:
    text = markdown_path.read_text(encoding="utf-8")
    parts = text.split("\n---\n")
    if len(parts) < 2:
        raise ValueError(f"cannot parse Marp slides from {markdown_path}")
    return parts[1:]


def renumber_image(image_blob: bytes, page_number: int) -> Image.Image:
    image = Image.open(BytesIO(image_blob)).convert("RGB")
    draw = ImageDraw.Draw(image)
    page_font = ImageFont.truetype(str(FONT_REGULAR), 42)
    draw.rectangle((2340, 1310, 2560, 1440), fill="white")
    draw.text((2520, 1394), str(page_number), font=page_font, fill=(112, 112, 112), anchor="rs")
    return image


def background_media_path(archive: zipfile.ZipFile, slide_number: int) -> str:
    relationships_path = f"ppt/slides/_rels/slide{slide_number}.xml.rels"
    root = ET.fromstring(archive.read(relationships_path))
    namespace = "{http://schemas.openxmlformats.org/package/2006/relationships}"
    for relationship in root.findall(f"{namespace}Relationship"):
        if relationship.attrib.get("Id") != "rId1":
            continue
        target = relationship.attrib["Target"]
        return posixpath.normpath(posixpath.join("ppt/slides", target))
    raise RuntimeError(f"slide {slide_number} has no rId1 background relationship")


def png_bytes(image: Image.Image) -> bytes:
    stream = BytesIO()
    image.save(stream, format="PNG", optimize=True)
    return stream.getvalue()


def replace_zip_entries(path: Path, replacements: dict[str, bytes]) -> None:
    with tempfile.NamedTemporaryFile(
        prefix=f".{path.stem}.media.", suffix=".pptx", dir=path.parent, delete=False
    ) as handle:
        temporary_path = Path(handle.name)
    try:
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary_path, "w") as target:
            existing = set(source.namelist())
            missing = set(replacements) - existing
            if missing:
                raise RuntimeError(f"saved PPTX lost background media: {sorted(missing)}")
            for info in source.infolist():
                target.writestr(info, replacements.get(info.filename, source.read(info.filename)))
        os.replace(temporary_path, path)
    finally:
        temporary_path.unlink(missing_ok=True)


def main() -> None:
    args = parse_args()
    for path in (args.pptx, args.markdown, args.poster, FONT_REGULAR, FONT_BOLD):
        if not path.is_file():
            raise FileNotFoundError(path)

    chunks = slide_chunks(args.markdown)
    if len(chunks) != 49:
        raise ValueError(f"expected 49 source slides after insertion, got {len(chunks)}")

    presentation = Presentation(str(args.pptx))
    if len(presentation.slides) != 48:
        raise ValueError(f"expected the verified 48-slide base deck, got {len(presentation.slides)}")

    # Marp stores each raster as the slide background rather than a picture
    # shape. Cache updated background assets before python-pptx rewrites the
    # package, then restore those exact media parts afterward.
    background_replacements: dict[str, bytes] = {}
    with zipfile.ZipFile(args.pptx, "r") as archive:
        for old_index in range(35, 49):
            new_index = old_index + 1
            chunk = chunks[new_index - 1]
            if "<!-- _paginate: false -->" in chunk:
                continue
            media_path = background_media_path(archive, old_index)
            background_replacements[media_path] = png_bytes(
                renumber_image(archive.read(media_path), new_index)
            )

    new_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    for shape in list(new_slide.shapes):
        shape._element.getparent().remove(shape._element)
    slide_image = render_slide(args.poster)
    stream = BytesIO()
    slide_image.save(stream, format="PNG", optimize=True)
    stream.seek(0)
    new_slide.shapes.add_picture(
        stream,
        0,
        0,
        width=presentation.slide_width,
        height=presentation.slide_height,
    )

    slide_id_list = presentation.slides._sldIdLst
    inserted_id = slide_id_list[-1]
    slide_id_list.remove(inserted_id)
    slide_id_list.insert(INSERT_SLIDE - 1, inserted_id)

    with tempfile.NamedTemporaryFile(
        prefix=f".{args.pptx.stem}.", suffix=".pptx", dir=args.pptx.parent, delete=False
    ) as handle:
        temporary_path = Path(handle.name)
    try:
        presentation.save(str(temporary_path))
        os.replace(temporary_path, args.pptx)
    finally:
        temporary_path.unlink(missing_ok=True)

    replace_zip_entries(args.pptx, background_replacements)

    print(f"inserted extreme funnel slide {INSERT_SLIDE}; deck now has 49 slides")


if __name__ == "__main__":
    main()
